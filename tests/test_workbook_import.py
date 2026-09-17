from io import BytesIO
import openpyxl
from rbe.workbook import import_toc_workbook

def make_fixture():
    wb=openpyxl.Workbook(); alignment=wb.active; alignment.title="PO and CO Alignment"; matrix=wb.create_sheet("Articulation Matrix"); co=wb.create_sheet("CO Calculation"); po=wb.create_sheet("PO Calculation")
    alignment["A2"]="Faculty"; alignment["E4"]="BCA"; alignment["E5"]="Computer Science"; alignment["E6"]="Theory of Computation"; alignment["E7"]="2026"; alignment["E8"]="Instructor"
    for i in range(1,5): alignment.cell(12,6+i).value=f"Outcome {i}"
    for c,v in enumerate(["CT 1","CT 2","Ass","ESE"],5): co.cell(10,c).value=v
    for c,v in enumerate([15,15,10,60],5): co.cell(11,c).value=v
    alloc=[[4,0,3,9],[3,4,2,16],[8,4,2,14],[0,7,3,21]]
    for r,row in enumerate(alloc,12):
        for c,v in enumerate(row,5): co.cell(r,c).value=v
    co.cell(20,1).value=1; co.cell(20,2).value=1; co.cell(20,3).value="A1"
    for c,v in enumerate([10,12,8,40],5): co.cell(20,c).value=v
    headers=["PO 1","PO 2","PO 3","PO 4","PO 5","PO 6","PO 7","PO 8","PO 9","PO 10","PO 11","PO 12","PSO-1","PSO-2","PSO-3"]
    for c,h in enumerate(headers,2): matrix.cell(7,c).value=h
    for r in range(8,12):
        matrix.cell(r,1).value=f"CO{r-7}"
        for c in range(2,17): matrix.cell(r,c).value=0
    matrix.cell(8,2).value=3
    for c,h in enumerate(headers[:15],4): po.cell(7,c).value=h
    po.cell(8,4).value=3; po.cell(15,11).value="#DIV/0!"
    bio=BytesIO(); wb.save(bio); bio.seek(0); return bio

def test_workbook_import_and_zero_denominator_warning():
    data=import_toc_workbook(make_fixture())
    assert data["metadata"]["course"]=="Theory of Computation"
    assert len(data["students"])==1
    assert data["mapping"]["CO1"]["PO 1"]==3
    assert data["po_attainment"]["PO 8"] is None
    assert any("#DIV/0!" in x for x in data["warnings"])
