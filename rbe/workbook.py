from io import BytesIO
from typing import Any
import openpyxl
from .obe import OBEStudent, calculate_co_results, weighted_po_attainment

REQUIRED_SHEETS = {"FLAT", "Matrix", "CO Calculation", "PO Calculation"}

def _load(source: Any, data_only=False):
    if hasattr(source, "read"):
        data = source.read()
        try:
            source.seek(0)
        except Exception:
            pass
        return openpyxl.load_workbook(BytesIO(data), data_only=data_only, read_only=False)
    return openpyxl.load_workbook(source, data_only=data_only, read_only=False)

def _txt(v):
    return "" if v is None else str(v).strip()

def import_toc_workbook(source: Any) -> dict:
    wb = _load(source, True)
    missing = REQUIRED_SHEETS - set(wb.sheetnames)
    if missing:
        raise ValueError("Workbook is missing required sheet(s): " + ", ".join(sorted(missing)))
    flat, co, matrix, poc = wb["FLAT"], wb["CO Calculation"], wb["Matrix"], wb["PO Calculation"]
    metadata = {"faculty": flat["A2"].value, "class": flat["E4"].value, "branch": flat["E5"].value, "course": flat["E6"].value, "year": flat["E7"].value, "instructor": flat["E8"].value}
    co_ids = [f"CO{i}" for i in range(1, 5)]
    descriptions = {f"CO{i}": flat.cell(12, 6 + i).value for i in range(1, 5)}
    components, max_marks = [], {}
    for c in range(5, 9):
        name = _txt(co.cell(10, c).value)
        if name:
            components.append(name)
            max_marks[name] = float(co.cell(11, c).value or 0)
    allocations = {cid: {components[k]: float(co.cell(row, 5 + k).value or 0) for k in range(len(components))} for row, cid in zip(range(12, 16), co_ids)}
    students = []
    for row in range(20, min(co.max_row, 1000) + 1):
        serial, roll, name = co.cell(row, 1).value, co.cell(row, 2).value, co.cell(row, 3).value
        if serial is None and roll is None and name is None:
            continue
        if roll is None:
            continue
        marks = {components[k]: float(co.cell(row, 5 + k).value or 0) for k in range(len(components))}
        students.append(OBEStudent(str(roll), _txt(name) or str(roll), marks))
    headers = [_txt(matrix.cell(7, c).value) for c in range(2, matrix.max_column + 1)]
    mapping = {cid: {headers[j]: float(matrix.cell(row, 2 + j).value or 0) for j in range(len(headers)) if headers[j]} for row, cid in zip(range(8, 12), co_ids)}
    student_rows, summary = calculate_co_results(students, max_marks, allocations)
    for cid in co_ids:
        summary[cid]["workbook_mean_score01"] = round(summary[cid]["mean_score01"], 2) if summary[cid]["mean_score01"] is not None else None
        summary[cid]["workbook_mean_level"] = round(summary[cid]["mean_level"], 2) if summary[cid]["mean_level"] is not None else None
    co_attainment = {cid: summary[cid]["workbook_mean_level"] for cid in co_ids}
    po_attainment = weighted_po_attainment(co_attainment, mapping)
    cached = {}
    for c in range(4, min(poc.max_column, 18) + 1):
        h = _txt(poc.cell(7, c).value)
        if h:
            cached[h] = poc.cell(15, c).value
    warnings = []
    for po, value in cached.items():
        if isinstance(value, str) and value.startswith("#"):
            warnings.append(f"Source workbook cached {po} attainment contains {value}; recomputation reports N/A when the mapping denominator is zero.")
    mismatch_found = False
    for row, cid in zip(range(8, 12), co_ids):
        for c, po in enumerate(headers, start=4):
            if not po:
                continue
            source_value, matrix_value = poc.cell(row, c).value, mapping[cid].get(po)
            if isinstance(source_value, (int, float)) and matrix_value is not None and abs(float(source_value) - float(matrix_value)) > 1e-9:
                warnings.append(f"Source workbook PO Calculation mapping differs from Matrix for {cid}/{po}; Matrix is used as the canonical mapping.")
                mismatch_found = True
                break
        if mismatch_found:
            break
    return {"metadata": metadata, "co_descriptions": descriptions, "components": components, "max_marks": max_marks, "co_allocations": allocations, "students": student_rows, "co_summary": summary, "mapping": mapping, "po_attainment": po_attainment, "source_po_cached": cached, "warnings": warnings}
